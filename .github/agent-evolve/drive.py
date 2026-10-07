"""One-off: an agent learns slides over seven 45-minute days (graduating from day 5 at the earliest), then sits an exam that may only use its MCP server."""
import base64
import io
import json
import os
from pathlib import Path
import re
import sys
import time

import requests

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)

from app.agent.session import AgentSessionManager  # noqa: E402
from app.agent import sandbox  # noqa: E402
from app.agent.tools import Workspace  # noqa: E402
from app.ai_providers import resolve_provider  # noqa: E402

MODE, OUT = sys.argv[1], Path(sys.argv[2]).resolve()
BASE, MODEL, KEY = os.environ["AGENT_BASE"], os.environ["AGENT_MODEL"], os.environ["AGENT_API_KEY"]
OUT.mkdir(parents=True, exist_ok=True)
HOME = Path.home()
DAY_MIN, DAYS, MIN_DAYS, EXAM_MIN, LEARN_CAP_MIN = 45, 7, 5, 75, 335
VI = re.compile(r"[ăâđêôơưáàảãạắằẳẵặấầẩẫậéèẻẽẹếềểễệíìỉĩịóòỏõọốồổỗộớờởỡợúùủũụứừửữựýỳỷỹỵ]", re.I)
URL = re.compile(r"https?://[^\s)\]>\"'`]+")
SCRATCH = re.compile(r"smoke|test|tmp|probe|scratch|draft", re.I)

RULES = """Bạn đang tham gia khóa TỰ HỌC 7 ngày để làm slide PowerPoint chuyên nghiệp như một designer nhiều năm kinh nghiệm.
LUẬT (bắt buộc):
1. Mỗi ngày = 45 phút làm việc thật, đồng hồ do hệ thống đếm. Bạn không tự kết thúc ngày được: xong mục tiêu sớm thì tự nâng độ khó (bộ slide mới, kiểu bố cục mới, sửa điểm yếu đã ghi, siết check_deck) cho tới khi hết giờ. Báo xong sớm hay ngồi chờ đều tính là bỏ học.
2. Tốt nghiệp chỉ được xét từ cuối ngày 5 trở đi (tối thiểu 5/7 ngày) và chỉ khi đạt ĐỦ 7 tiêu chí A–G dưới đây. Hệ thống chấm bằng script, không tin lời tự khai. Muốn xin tốt nghiệp thì ghi đúng dòng "XIN TỐT NGHIỆP" trong báo cáo cuối ngày.
3. Khi nhận [Time notice], việc đầu tiên là cập nhật mục "## Ngày N" trong deck-lab/JOURNAL.md (mục tiêu, nguồn đã đọc kèm URL, đã làm, số lỗi check_deck trước/sau, bài học, kế hoạch ngày mai), rồi mới viết báo cáo ngắn.
4. Không bao giờ báo "xong" hay "sạch" khi chưa có bằng chứng: output lệnh, báo cáo check_deck và ảnh đã xem bằng view_image. Lỗi nào gặp cũng phải ghi thật vào nhật ký.
5. Nội dung tiếng Việt phải có dấu đầy đủ, dùng font có đủ glyph tiếng Việt (ví dụ DejaVu Sans) và kiểm tra bằng ảnh render. Không được bỏ dấu để né lỗi.
6. Không làm trước đề thi: trong khóa học không làm slide về kiến trúc repo manga-translator.
7. Server MCP mới khai báo trong .mcp.json được hệ thống tin cậy ở lần nghỉ kế tiếp (sau khi bạn gọi goal_done, hoặc đầu ngày sau); khi đó các tool mcp__deck__* xuất hiện. Khi có nhiều tool MCP, chúng nằm sau tool_search: gọi tool_search với từ khóa "deck" để nạp trước khi dùng.
CÔNG CỤ: mcp__tinyfish__search và mcp__tinyfish__fetch để nghiên cứu; python-pptx đã cài; xem kết quả bằng
`soffice -env:UserInstallation=file:///tmp/lo_profile --headless --convert-to pdf --outdir OUT FILE.pptx` rồi `pdftoppm -png -r 50 FILE.pdf PREFIX`, sau đó view_image từng trang.
TIÊU CHÍ ĐẦU RA (script chấm tự động mỗi cuối ngày):
A. deck-lab/RESEARCH.md: quy tắc của dân chuyên nghiệp bằng con số cụ thể (lưới, lề, cỡ chữ theo pt, số chữ mỗi slide, tỉ lệ tương phản), trích ít nhất 8 URL khác nhau đã đọc bằng mcp__tinyfish__fetch.
B. .agents/skills/pptx/SKILL.md: front matter có name và description; ít nhất 8 mục "## " (bố cục và lưới, thang cỡ chữ, màu theo vai trò kèm mã hex và tỉ lệ tương phản, ít nhất 8 kiểu slide mẫu, biểu đồ, chuyển trang và animation, tiếng Việt, quy trình tự kiểm tra); trích ít nhất 5 URL; dài ít nhất 6000 ký tự.
C. Plugin của bạn (plugin_write, dùng tham số path cho file dài) cung cấp tool make_deck và check_deck, mỗi tool có JSON schema parameters đầy đủ (properties, required, mô tả từng trường). Lỗi trả về phải nói rõ trường nào sai và vì sao, không được chỉ là tên một key. check_deck phải bắt được: chữ tràn khung, khung chồng nhau, shape ra ngoài slide, chữ dưới 12pt, tương phản dưới 4.5:1, khung chữ trống, slide quá 70 chữ, thiếu chuyển trang. Hệ thống sẽ đưa một file .pptx cố tình có lỗi cho check_deck của bạn.
D. .mcp.json ở thư mục làm việc khai báo server tên "deck" phục vụ make_deck và check_deck. Repo này có sẵn lệnh `python run.py mcp --folder . --write` (run.py ở ngay thư mục gốc) phục vụ tool của mọi plugin qua MCP; nếu tự viết server thì stdout chỉ được chứa JSON-RPC. Lưu bản ghi phiên test thật (initialize, tools/list, tools/call tạo file .pptx) vào deck-lab/mcp_test.log.
E. Portfolio: ít nhất 8 bộ .pptx trong deck-lab/ (tên file không chứa test, smoke, tmp, probe, draft), mỗi bộ ít nhất 6 slide, phủ ít nhất 6 kiểu bố cục (bìa, mục lục, thẻ, so sánh hai cột, quy trình, biểu đồ số liệu, trích dẫn hoặc kết luận, sơ đồ); ít nhất 2 bộ tiếng Việt có dấu; mọi bộ: không shape ra ngoài slide, không chữ dưới 12pt, không khung chữ trống, không slide quá 70 chữ, có chuyển trang.
F. deck-lab/JOURNAL.md có mục "## Ngày N" cho mọi ngày đã học, mỗi mục ít nhất 120 từ.
G. Thi thử: ít nhất 1 bộ (ví dụ kiến trúc thư viện python-pptx đã cài) làm hoàn toàn bằng các tool mcp__deck__* trong phiên này, ghi lại trong JOURNAL.md với chữ "thi thử"."""
PLAN = {
    1: "Nghiên cứu: dùng mcp__tinyfish__search và mcp__tinyfish__fetch đọc ít nhất 8 nguồn chuyên nghiệp về thiết kế slide (bố cục, lưới, typography, màu, biểu đồ, kể chuyện bằng dữ liệu) và viết deck-lab/RESEARCH.md. Làm bộ slide đầu tiên bằng tay theo nghiên cứu, render, xem, sửa.",
    2: "Viết SKILL.md v1 và script tự kiểm tra trong .agents/skills/pptx/ từ nghiên cứu và lỗi ngày 1. Làm 2 bộ mới chỉ theo skill (1 bộ tiếng Việt có dấu), nâng cấp skill sau mỗi bộ.",
    3: "Biến skill thành plugin bằng plugin_write với make_deck và check_deck đạt tiêu chí C. Làm 3 bộ chỉ bằng plugin, ba kiểu khác nhau (số liệu có biểu đồ, quy trình từng bước, so sánh hai cột).",
    4: "Tiến hóa plugin ít nhất 2 vòng: tìm điểm yếu về bố cục, hiệu ứng, màu, chữ, cỡ chữ từ ảnh render; viết lại plugin; dựng lại các bộ; ghi deck-lab/EVOLUTION.md với số đo trước/sau từng vòng.",
    5: "Đưa plugin lên MCP đạt tiêu chí D và test thật; gọi goal_done để hệ thống tin cậy server, rồi thi thử đạt tiêu chí G. Đối chiếu toàn bộ tiêu chí A–G; đủ thì xin tốt nghiệp.",
    6: "Ngày bổ sung: sửa mọi tiêu chí chưa đạt, nâng chất lượng portfolio, thi thử lại.",
    7: "Ngày cuối: hoàn thiện mọi tiêu chí còn thiếu, thi thử lần cuối, viết tổng kết khóa học trong JOURNAL.md.",
}
KEEP_GOING = ("/goal [Đồng hồ] Ngày {n} còn khoảng {m} phút nên chưa được dừng. {note}Tự nâng độ khó: bộ slide mới, kiểu bố cục mới, "
              "sửa điểm yếu trong nhật ký, siết check_deck, hoặc các tiêu chí còn thiếu: {fails}")
USER_TEXT = ("repo manga-translator có những thư mục file nào kiến trúc tổng quát sao và đi sâu vào từng lớp kiến trúc "
             "sau đó tạo slide pptx và xuất thành bài cuối cùng")
EXAM = ("/goal " + USER_TEXT + """

YÊU CẦU ĐẦU RA BẮT BUỘC (chấm tự động bằng script):
1. Tạo slide CHỈ bằng tool mcp__deck__make_deck và kiểm tra bằng mcp__deck__check_deck (nếu chưa thấy chúng trong danh sách tool thì gọi tool_search với từ khóa "deck" để nạp); không tự viết code python-pptx hay script tạo slide riêng.
2. Đọc code thật trước khi viết: mọi đường dẫn file hay thư mục ghi trên slide phải tồn tại trong repo.
3. 12–16 slide, tiếng Việt có dấu đầy đủ: bìa; cây thư mục cấp cao nhất kèm vai trò từng thư mục; sơ đồ kiến trúc tổng quát các lớp; mỗi lớp kiến trúc ít nhất 1 slide đi sâu (thành phần chính, file tiêu biểu, dữ liệu vào và ra); luồng xử lý một chương từ tải về đến xuất file; điểm mạnh và rủi ro; kết luận.
4. Không khung chữ trống, không chữ dưới 12pt, không shape ra ngoài slide, không slide quá 70 chữ; mcp__deck__check_deck phải báo 0 lỗi trên file cuối.
5. Render ra PNG và xem từng trang bằng view_image, sửa tới khi đẹp.
6. Xuất final/manga-translator-architecture.pptx và final/manga-translator-architecture.pdf.
7. Báo cáo cuối bằng tiếng Việt: dàn ý từng slide và bằng chứng (output check_deck, các trang đã xem).""")


def probe_vision() -> bool:
    """Whether the model reads an image: ask for the colour of a small red square."""
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (64, 64), (220, 20, 20)).save(buf, format="PNG")
    url = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    body = {"model": MODEL, "max_tokens": 20, "messages": [{"role": "user", "content": [
        {"type": "text", "text": "What colour is this image? Answer with one word."}, {"type": "image_url", "image_url": {"url": url}}]}]}
    try:
        reply = requests.post(f"{BASE}/chat/completions", json=body, headers={"Authorization": f"Bearer {KEY}"}, timeout=120)
        text = reply.json()["choices"][0]["message"]["content"] if reply.ok else reply.text[:200]
    except Exception as exc:
        text = f"{type(exc).__name__}: {exc}"
    print("vision probe:", str(text)[:200], flush=True)
    return "red" in str(text).lower()


def setup() -> tuple:
    vision = probe_vision()
    folder = HOME / ".manga-agent"
    folder.mkdir(parents=True, exist_ok=True)
    profile = folder / "profile.json"
    data = json.loads(profile.read_text()) if profile.is_file() else {}
    profile.write_text(json.dumps({**data, "token_budget": 10 ** 12, "max_steps": 5000, "goal_turns": 300, "vision": vision}))
    # TinyFish search and fetch as a user-scope MCP server, so every session sees mcp__tinyfish__* tools.
    server = {"command": sys.executable, "args": [str(Path(__file__).with_name("tinyfish_mcp.py"))]}
    (HOME / ".claude.json").write_text(json.dumps({"mcpServers": {"tinyfish": server}}))
    label = os.environ.get("AGENT_LABEL", "Provider")
    provider = resolve_provider(label.lower(), label=label, protocol="openai", api_base=BASE)
    manager = AgentSessionManager(OUT.parent / f"store-{MODE}", home=HOME)
    return vision, provider, manager, sandbox.Policy("workspace-write", True)


def trust_workspace_mcp(session) -> list[str]:
    """Approve the workspace's MCP servers as a watching person would; only called while the session is idle."""
    from app.agent import mcp

    session._ensure_mcp()
    # A server added or changed in .mcp.json after the session began is only read again on /plugins reload.
    rows = [r for r in mcp.configured(session.workspace.root, session.home) if r["scope"] == "workspace"]
    if any(r["name"] not in session.mcp_status or session.mcp_status[r["name"]]["digest"] != mcp.config_hash(r["config"]) for r in rows):
        session.command("/plugins reload")
    names = []
    for name, status in list(session.mcp_status.items()):
        if status.get("scope") == "workspace" and status.get("state") not in ("running",):
            try:
                session.trust_mcp(name)
                names.append(name)
            except Exception as exc:
                print(f"trust {name} failed: {exc}", flush=True)
    return names


def watch(session, label: str, seconds: float, sink: list) -> str:
    """Wait for the turn to end, approving what asks; returns the last assistant text."""
    start, seen = time.time(), len(session.events)
    while True:
        time.sleep(3)
        for event in session.events[seen:]:
            seen += 1
            sink.append(event)
            if event["type"] == "tool":
                print(f"[{label} {int(time.time() - start)}s] {event['name']} {'ok' if event.get('ok') else 'FAILED'}: {str(event.get('output', ''))[:140]!r}", flush=True)
            elif event["type"] in ("notice", "error"):
                print(f"[{label}] {event['type']}: {event.get('text', '')[:300]}", flush=True)
        if session.status == "waiting" and session.pending:
            print(f"[{label}] approving {session.pending['name']}", flush=True)
            session.decide("allow")
        if session.status == "idle" and time.time() - start > 5:
            break
        if time.time() - start > seconds:
            print(f"[{label}] over the hard limit; stopping", flush=True)
            session.stop()
            time.sleep(20)
            break
    return next((e["text"] for e in reversed(sink) if e["type"] == "assistant" and e.get("text")), "")


# Grading.

def read(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except OSError:
        return ""


def lint_deck(path: Path) -> dict:
    """Problems a strict reviewer would flag, measured from the file itself."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    import zipfile

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        return {"slides": 0, "problems": [f"cannot open: {exc}"], "text": "", "transitions": 0}
    width, height, problems, texts = prs.slide_width, prs.slide_height, [], []
    for number, slide in enumerate(prs.slides, 1):
        words = 0
        for shape in slide.shapes:
            box = [shape.left, shape.top, shape.width, shape.height]
            if None not in box and (box[0] < -9525 or box[1] < -9525 or box[0] + box[2] > width + 9525 or box[1] + box[3] > height + 9525):
                problems.append(f"slide {number}: '{shape.name}' goes off the slide")
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            texts.append(text)
            words += len(text.split())
            if not text and (shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX or shape.is_placeholder):
                problems.append(f"slide {number}: empty text box '{shape.name}'")
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size is not None and run.font.size.pt < 12:
                        problems.append(f"slide {number}: {run.font.size.pt:g}pt text '{run.text.strip()[:30]}'")
        if words > 70:
            problems.append(f"slide {number}: {words} words (over 70)")
    with zipfile.ZipFile(path) as archive:
        transitions = sum(1 for name in archive.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", name) and b"<p:transition" in archive.read(name))
    return {"slides": len(prs.slides), "problems": problems, "text": "\n".join(texts), "transitions": transitions}


def bad_deck(path: Path) -> None:
    """A deck with planted faults for the agent's check_deck to find."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tiny = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1.2), Inches(0.3)).text_frame
    tiny.text = "Chữ rất nhỏ và quá dài để vừa một khung hẹp như thế này, lặp lại nhiều lần cho chắc chắn tràn khung " * 3
    for paragraph in tiny.paragraphs:
        for run in paragraph.runs:
            run.font.size, run.font.color.rgb = Pt(8), RGBColor(0xEE, 0xEE, 0xEE)
    slide.shapes.add_textbox(Inches(12.5), Inches(6.8), Inches(3), Inches(2)).text_frame.text = "Ra ngoài slide"
    slide.shapes.add_textbox(Inches(4), Inches(3), Inches(3), Inches(1))
    slide.shapes.add_textbox(Inches(1.1), Inches(1.1), Inches(3), Inches(1)).text_frame.text = "Chồng lên khung khác"
    prs.save(str(path))


def grade_learning(session, days_done: int) -> dict:
    """Criteria A–G, checked from files and from the agent's own tools."""
    result = {}
    research = read(ROOT / "deck-lab" / "RESEARCH.md")
    urls = set(URL.findall(research))
    result["A"] = (len(urls) >= 8, f"RESEARCH.md có {len(urls)} URL khác nhau (cần ≥8)")
    skill = read(ROOT / ".agents" / "skills" / "pptx" / "SKILL.md")
    head = skill.split("---")[1] if skill.startswith("---") and skill.count("---") >= 2 else ""
    sections = len(re.findall(r"^## ", skill, re.M))
    skill_urls = len(set(URL.findall(skill)))
    ok_b = "name:" in head and "description:" in head and sections >= 8 and skill_urls >= 5 and len(skill) >= 6000
    result["B"] = (ok_b, f"SKILL.md: front matter {'đủ' if 'name:' in head and 'description:' in head else 'thiếu'}, {sections} mục ##, {skill_urls} URL, {len(skill)} ký tự")
    tools = session.registry.tools
    notes, ok_c = [], True
    for name in ("make_deck", "check_deck"):
        params = tools[name].spec.get("parameters") if name in tools else None
        if not params or not params.get("properties") or not params.get("required"):
            ok_c = False
            notes.append(f"{name}: {'không có' if name not in tools else 'schema thiếu properties hoặc required'}")
    if "check_deck" in tools:
        probe = ROOT / "deck-lab" / "grader_probe_bad.pptx"
        probe.parent.mkdir(parents=True, exist_ok=True)
        bad_deck(probe)
        props = tools["check_deck"].spec["parameters"].get("properties") or {}
        key = next((k for k in props if "path" in k.lower() or "file" in k.lower()), (tools["check_deck"].spec["parameters"].get("required") or ["path"])[0])
        output, ok = session._run_call({"id": "grader-c", "name": "check_deck", "args": {key: str(probe)}})
        probe.unlink(missing_ok=True)
        clean = re.search(r"\bclean\b|no problems|\"problems\": \[\]|\b0 (issues?|problems?|lỗi)", output, re.I)
        if clean or len(output) < 40:
            ok_c = False
        notes.append(f"check_deck trên file cài lỗi: {output[:300]!r}")
    result["C"] = (ok_c, "; ".join(notes))
    config = read(ROOT / ".mcp.json")
    log = read(ROOT / "deck-lab" / "mcp_test.log")
    live = {"mcp__deck__make_deck", "mcp__deck__check_deck"} <= set(session.mcp_tools)
    ok_d = '"deck"' in config and "tools/list" in log and "tools/call" in log and live
    result["D"] = (ok_d, f".mcp.json {'có' if chr(34) + 'deck' + chr(34) in config else 'không có'} deck, mcp_test.log {'đủ' if 'tools/call' in log else 'thiếu'}, tool mcp__deck__ {'chạy' if live else 'không chạy'}")
    decks = [p for p in sorted((ROOT / "deck-lab").rglob("*.pptx")) if not SCRATCH.search(p.name)]
    bad, vietnamese = [], 0
    for deck in decks:
        lint = lint_deck(deck)
        if len(VI.findall(lint["text"])) >= 30:
            vietnamese += 1
        if lint["slides"] < 6 or lint["problems"] or lint["transitions"] == 0:
            bad.append(f"{deck.name}: {lint['slides']} slide, {lint['transitions']} chuyển trang, {len(lint['problems'])} lỗi {lint['problems'][:2]}")
    ok_e = len(decks) >= 8 and not bad and vietnamese >= 2
    result["E"] = (ok_e, f"{len(decks)} bộ (cần ≥8), {vietnamese} bộ tiếng Việt (cần ≥2), bộ chưa đạt: {bad[:4]}")
    journal = read(ROOT / "deck-lab" / "JOURNAL.md")
    short = []
    for day in range(1, days_done + 1):
        match = re.search(rf"^## Ngày {day}\b(.*?)(?=^## |\Z)", journal, re.M | re.S)
        if not match or len(match.group(1).split()) < 120:
            short.append(day)
    result["F"] = (not short, f"ngày thiếu hoặc dưới 120 từ: {short}")
    mcp_made = sum(1 for e in session.events if e["type"] == "tool" and e["name"] == "mcp__deck__make_deck" and e.get("ok"))
    result["G"] = (mcp_made >= 1 and "thi thử" in journal.lower(), f"{mcp_made} lần mcp__deck__make_deck thành công, nhật ký {'có' if 'thi thử' in journal.lower() else 'chưa có'} thi thử")
    return {k: {"ok": v[0], "detail": v[1]} for k, v in result.items()}


def learn() -> None:
    vision, provider, manager, policy = setup()
    report = {"model": MODEL, "vision": vision, "sandbox": sandbox.backend(), "days": []}
    session = manager.create(provider, KEY, MODEL, Workspace(ROOT, policy), "auto")
    session._ensure_mcp()
    report["mcp_at_start"] = {"status": session.command("/mcp")["message"], "tools": sorted(session.mcp_tools)}
    print("mcp at start:", json.dumps(report["mcp_at_start"], ensure_ascii=False), flush=True)
    began, grade, graduated = time.time(), None, False
    for n in range(1, DAYS + 1):
        if (time.time() - began) / 60 > LEARN_CAP_MIN - DAY_MIN:
            print("no time left for another full day", flush=True)
            break
        trusted = trust_workspace_mcp(session)
        fails = [f"{k}: {v['detail']}" for k, v in (grade or {}).items() if not v["ok"]]
        status = ""
        if grade is not None:
            status = f"Chấm tự động cuối ngày {n - 1}: đạt {sum(v['ok'] for v in grade.values())}/7 tiêu chí. Chưa đạt: {' | '.join(fails) or 'không'}. "
        if trusted:
            status += f"Hệ thống đã tin cậy server MCP {', '.join(trusted)}; các tool mcp__{trusted[0]}__* đã dùng được. "
        prompt = f"/goal Ngày {n}/7 bắt đầu (45 phút). {status}Mục tiêu hôm nay: {PLAN[n]} Làm xong sớm thì tự nâng độ khó cho tới khi hết giờ."
        if n == 1:
            prompt = "/goal " + RULES + "\n\n" + prompt[len("/goal "):]
        day_start, events, last = time.time(), [], ""
        session.set_deadline(DAY_MIN * 60)
        session.send(prompt)
        while True:
            left = DAY_MIN * 60 - (time.time() - day_start)
            last = watch(session, f"day{n}", left + 180, events) or last
            left = DAY_MIN * 60 - (time.time() - day_start)
            if left < 240:
                break
            note = ""
            names = trust_workspace_mcp(session)
            if names:
                note = f"Hệ thống vừa tin cậy server MCP {', '.join(names)}; các tool mcp__{names[0]}__* đã dùng được. "
            if "XIN TỐT NGHIỆP" in last and n < MIN_DAYS:
                note += f"Đơn xin tốt nghiệp bị bác: mới ngày {n}, phải học tối thiểu {MIN_DAYS} ngày. "
            print(f"[day{n}] idle with {int(left / 60)} min left; keep going", flush=True)
            session.set_deadline(left)
            session.send(KEEP_GOING.format(n=n, m=int(left / 60), note=note, fails=" | ".join(fails) or "tự tìm điểm yếu"))
        with open(OUT / f"day{n}-events.jsonl", "w", encoding="utf-8") as handle:
            for event in events:
                handle.write(json.dumps(event, ensure_ascii=False, default=str) + "\n")
        trust_workspace_mcp(session)
        grade = grade_learning(session, n)
        tools = [e["name"] for e in events if e["type"] == "tool"]
        asked = "XIN TỐT NGHIỆP" in last
        day = {"day": n, "minutes": round((time.time() - day_start) / 60, 1), "tool_counts": {t: tools.count(t) for t in sorted(set(tools))},
               "failed_tools": sum(1 for e in events if e["type"] == "tool" and e.get("ok") is False), "usage": dict(session.usage),
               "rate_limited": dict(session.rate_limited), "asked_to_graduate": asked, "grade": grade, "report": last[:6000]}
        report["days"].append(day)
        passed = sum(v["ok"] for v in grade.values())
        print(f"=== day {n}: {passed}/7 criteria; asked to graduate: {asked}; {json.dumps(grade, ensure_ascii=False)[:1500]}", flush=True)
        if asked and n >= MIN_DAYS and passed == 7:
            graduated = True
        report["graduated"], report["days_studied"] = graduated, n
        (OUT / "summary.json").write_text(json.dumps(report, ensure_ascii=False, indent=1))
        if graduated:
            print(f"graduated after day {n}", flush=True)
            break
    report["plugins"] = session.command("/plugins")["message"]
    report["mcp"] = session.command("/mcp")["message"]
    from app.agent import client
    report["rate_limits_total"] = client.RATE_LIMITS
    (OUT / "summary.json").write_text(json.dumps(report, ensure_ascii=False, indent=1))
    manager.close_all()


def grade_exam(session, events: list) -> dict:
    final = ROOT / "final" / "manga-translator-architecture.pptx"
    result = {"pptx": final.is_file(), "pdf": (ROOT / "final" / "manga-translator-architecture.pdf").is_file()}
    lint = lint_deck(final) if final.is_file() else {"slides": 0, "problems": ["missing"], "text": "", "transitions": 0}
    slides_with_vi = 0
    if final.is_file():
        from pptx import Presentation

        for slide in Presentation(str(final)).slides:
            text = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
            slides_with_vi += bool(VI.search(text))
    cited = sorted(set(re.findall(r"\b(?:app|gateway|tests|scripts|extension|deploy|docs|installer|models|data)(?:/[\w.\-]+)+/?", lint["text"])))
    missing = [p for p in cited if not (ROOT / p.rstrip("/")).exists()]
    calls = [e for e in events if e["type"] == "tool"]
    made = sum(1 for e in calls if e["name"] == "mcp__deck__make_deck" and e.get("ok"))
    checked = sum(1 for e in calls if e["name"] == "mcp__deck__check_deck")
    scripted = []
    for event in events:
        for call in event.get("calls") or []:
            blob = json.dumps(call.get("args") or {}, ensure_ascii=False)
            if call["name"] in ("write_file", "edit_file", "apply_patch", "run_command", "run_script") and re.search(r"from pptx|import pptx|Presentation\(", blob):
                scripted.append(call["name"])
    final_check = ""
    if final.is_file() and "mcp__deck__check_deck" in session.mcp_tools:
        schema = session.mcp_tools["mcp__deck__check_deck"][1].get("inputSchema") or {}
        props = schema.get("properties") or {}
        key = next((k for k in props if "path" in k.lower() or "file" in k.lower()), (schema.get("required") or ["path"])[0])
        session.mcp_loaded.add("mcp__deck__check_deck")
        final_check = session._run_call({"id": "grader-exam", "name": "mcp__deck__check_deck", "args": {key: str(final)}})[0][:1500]
    checks = {
        "files": result["pptx"] and result["pdf"],
        "slide_count_12_16": 12 <= lint["slides"] <= 16,
        "vietnamese_on_90pct_slides": lint["slides"] > 0 and slides_with_vi >= 0.9 * lint["slides"],
        "no_lint_problems": not lint["problems"],
        "cited_paths_exist": bool(cited) and not missing,
        "made_with_mcp": made >= 1,
        "checked_with_mcp": checked >= 1,
        "no_own_pptx_script": not scripted,
    }
    return {"checks": checks, "passed": sum(checks.values()), "total": len(checks), "slides": lint["slides"], "slides_with_vietnamese": slides_with_vi,
            "lint_problems": lint["problems"][:30], "cited_paths": cited, "missing_paths": missing, "mcp_make_calls": made,
            "mcp_check_calls": checked, "own_pptx_script_calls": scripted, "agent_check_deck_on_final": final_check}


def exam() -> None:
    vision, provider, manager, policy = setup()
    session = manager.create(provider, KEY, MODEL, Workspace(ROOT, policy), "auto")
    # The slide tools exist only through MCP: the agent's plugin rows are switched off and its .mcp.json servers trusted.
    for row in [r.id for r in session.kernel.rows.values() if r.source in ("agent", "user")]:
        session.command(f"/plugins disable {row}")
    trust_workspace_mcp(session)
    report = {"model": MODEL, "vision": vision, "setup": {"plugins": session.command("/plugins")["message"], "mcp": session.command("/mcp")["message"],
                                                          "mcp_tools": sorted(session.mcp_tools)}}
    print("exam setup:", json.dumps(report["setup"], ensure_ascii=False)[:2000], flush=True)
    events, start = [], time.time()
    session.set_deadline(EXAM_MIN * 60)
    session.send(EXAM)
    last = watch(session, "exam", EXAM_MIN * 60 + 180, events)
    with open(OUT / "exam-events.jsonl", "w", encoding="utf-8") as handle:
        for event in events:
            handle.write(json.dumps(event, ensure_ascii=False, default=str) + "\n")
    tools = [e["name"] for e in events if e["type"] == "tool"]
    report["exam"] = {"minutes": round((time.time() - start) / 60, 1), "tool_counts": {t: tools.count(t) for t in sorted(set(tools))},
                      "usage": dict(session.usage), "rate_limited": dict(session.rate_limited), "report": last[:8000]}
    report["grade"] = grade_exam(session, events)
    print("exam grade:", json.dumps(report["grade"], ensure_ascii=False)[:3000], flush=True)
    (OUT / "summary.json").write_text(json.dumps(report, ensure_ascii=False, indent=1))
    manager.close_all()


if __name__ == "__main__":
    learn() if MODE == "learn" else exam()

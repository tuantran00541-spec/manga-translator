"""pptx deck builder + QC plugin (v1)."""
import json
import os

EMU = 914400
SW, SH = 13.333, 7.5
DARK = {
    "bg": (0x14, 0x12, 0x2B), "card": (0x23, 0x1F, 0x42),
    "ink": (0xF4, 0xF1, 0xFF), "muted": (0xB8, 0xAE, 0xE6),
    "accent": (0x9F, 0x7A, 0xF0), "accent2": (0x5E, 0xE6, 0xC8),
}
LIGHT = {
    "bg": (0xF6, 0xF4, 0xFB), "card": (0xFF, 0xFF, 0xFF),
    "ink": (0x24, 0x1E, 0x3C), "muted": (0x5A, 0x53, 0x78),
    "accent": (0x6E, 0x46, 0xE0), "accent2": (0x0A, 0x6B, 0x55),
}

def _rgb(pal, key):
    from pptx.dml.color import RGBColor
    c = pal.get(key, DARK[key])
    return RGBColor(c[0], c[1], c[2])

def _bg(slide, pal):
    from pptx.util import Pt
    r = slide.shapes.add_shape(1, 0, 0, int(SW * EMU), int(SH * EMU))
    r.fill.solid(); r.fill.fore_color.rgb = _rgb(pal, "bg"); r.line.fill.background()
    return r

_TRANS = {"push": "push", "fade": "fade", "wipe": "wipe", "blinds": "blinds", "dissolve": "dissolve", "none": None}

def _slide_in(slide, name):
    """Insert <p:transition> after cSld/clrMapOvr (schema order), else LibreOffice drops it."""
    if _TRANS.get(name) is None:
        return
    from lxml import etree
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    el = etree.fromstring(f'<p:transition xmlns:p="{P}" spd="med"><p:{_TRANS[name]}/></p:transition>')
    slide._element.insert(2, el)

def _card(slide, x, y, w, h, pal, rounded=False, border=None):
    from pptx.util import Pt
    shape = slide.shapes.add_shape(5 if rounded else 1, int(x * EMU), int(y * EMU), int(w * EMU), int(h * EMU))
    if rounded:
        shape.adjustments[0] = 0.04
    shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(pal, "card")
    if border:
        shape.line.color.rgb = _rgb(pal, border); shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

def _text(slide, x, y, w, h, lines, space_after=6, align="left"):
    """lines: [{text, pt, color, bold}] color is role name or hex str."""
    from pptx.util import Pt
    from pptx.enum.text import PP_ALIGN
    tb = slide.shapes.add_textbox(int(x * EMU), int(y * EMU), int(w * EMU), int(h * EMU))
    tf = tb.text_frame; tf.word_wrap = True
    pal_map = {"ink": "ink", "muted": "muted", "accent": "accent", "accent2": "accent2"}
    global _PAL
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        r = p.add_run(); r.text = ln["text"]
        r.font.size = Pt(ln.get("pt", 14))
        r.font.name = "DejaVu Sans"
        col = ln.get("color", "ink")
        if col in pal_map:
            r.font.color.rgb = _rgb(_PAL, col)
        else:
            r.font.color.rgb = _hex(col)
        r.font.bold = bool(ln.get("bold", False))
    return tb

def _hex(s):
    from pptx.dml.color import RGBColor
    s = s.lstrip("#")
    return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))

def _add_chart(slide, x, y, w, h, data, pal):
    """data: {categories: [..], series: [{name, values:[..]}]}"""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    cd = CategoryChartData()
    cd.categories = data["categories"]
    for s in data["series"]:
        cd.add_series(s["name"], s["values"])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, int(x * EMU), int(y * EMU), int(w * EMU), int(h * EMU), cd)
    ch = gf.chart; ch.has_legend = True
    plot = ch.plots[0]
    for i, s in enumerate(data["series"]):
        plot.series[i].format.fill.solid()
        plot.series[i].format.fill.fore_color.rgb = _rgb(pal, s.get("color", "accent" if i == 0 else "muted"))
    return gf

# ---- layout builders ----

def L_title(s, data, pal):
    x0, y0 = 0.8, 1.4
    lines = [{"text": data["title"], "pt": 42, "color": "ink", "bold": True}]
    if data.get("subtitle"):
        lines.append({"text": data["subtitle"], "pt": 42, "color": "accent", "bold": True})
    _text(s, x0, y0, 11.7, 1.8, lines, space_after=0)
    if data.get("desc"):
        _text(s, 0.8, 3.8, 11.7, 1.0, [{"text": data["desc"], "pt": 18, "color": "muted"}], space_after=0)
    if data.get("footer"):
        _text(s, 0.8, 6.4, 11.7, 0.6, [{"text": data["footer"], "pt": 13, "color": "accent2", "bold": True}], space_after=0)

def L_bullets(s, data, pal):
    n = len(data["items"])
    if n <= 2:
        _card(s, 0.5, 0.5, 12.3, 6.5, pal)
        per = 4.6 / n
        y = 1.9
        for it in data["items"]:
            lines = []
            if it.get("label"):
                lines.append({"text": it["label"], "pt": 15, "color": "accent2", "bold": True})
            lines.append({"text": it["text"], "pt": 15, "color": "ink"})
            _text(s, 0.9, y, 11.5, per, lines, space_after=4)
            y += per
    else:
        _text(s, 0.5, 0.35, 12.3, 0.8, [{"text": data.get("title", ""), "pt": 26, "color": "accent", "bold": True}], space_after=0)
        half = (n + 1) // 2
        for col, chunk in enumerate([data["items"][:half], data["items"][half:]]):
            x = 0.5 + col * 6.3
            _card(s, x, 1.3, 5.9, 5.5, pal)
            per = 5.5 / len(chunk) if chunk else 1
            y = 1.55
            for it in chunk:
                lines = []
                if it.get("label"):
                    lines.append({"text": it["label"], "pt": 15, "color": "accent2", "bold": True})
                lines.append({"text": it["text"], "pt": 14, "color": "ink"})
                _text(s, x + 0.2, y, 5.5, per, lines, space_after=4)
                y += per

def L_cards(s, data, pal):
    n = len(data["cards"])
    gap = 0.2
    w = (12.3 - gap * (n - 1)) / n
    _text(s, 0.5, 0.35, 12.3, 0.8, [{"text": data.get("title", ""), "pt": 26, "color": "accent", "bold": True}], space_after=0)
    x = 0.5
    for c in data["cards"]:
        _card(s, x, 1.3, w, 5.5, pal)
        lines = []
        if c.get("kicker"):
            lines.append({"text": c["kicker"], "pt": 13, "color": "accent2", "bold": True})
        if c.get("title"):
            lines.append({"text": c["title"], "pt": 15, "color": "ink", "bold": True})
        lines.append({"text": c.get("text", ""), "pt": 13, "color": "muted"})
        if c.get("takeaway"):
            lines.append({"text": "★ " + c["takeaway"], "pt": 13, "color": "accent", "bold": True})
        _text(s, x + 0.15, 1.55, w - 0.3, 4.5, lines, space_after=4)
        x += w + gap

def L_two_panel(s, data, pal):
    for side, x0 in (("left", 0.5), ("right", 6.8)):
        d = data.get(side, {})
        _card(s, x0, 0.5, 6.0, 6.5, pal)
        lines = [{"text": d.get("title", ""), "pt": 20, "color": "accent", "bold": True}]
        for p in d.get("paragraphs", []):
            lines.append({"text": p, "pt": 15, "color": "ink"})
        _text(s, x0 + 0.3, 0.9, 5.4, 5.5, lines, space_after=6)

def L_chart(s, data, pal):
    _text(s, 0.5, 0.35, 12.3, 0.8, [{"text": data.get("title", ""), "pt": 26, "color": "accent", "bold": True}], space_after=0)
    d = data["data"]
    _add_chart(s, 0.8, 1.4, 7.5, 5.5, d, pal)
    note_lines = []
    if data.get("note_title"):
        note_lines.append({"text": data["note_title"], "pt": 15, "color": "accent2", "bold": True})
    for n in data.get("note", []):
        note_lines.append({"text": n, "pt": 15, "color": "ink"})
    if note_lines:
        _text(s, 9.0, 1.8, 3.8, 4.5, note_lines, space_after=6)

def L_numbered_rows(s, data, pal):
    _card(s, 0.5, 0.5, 12.3, 6.5, pal)
    _text(s, 0.9, 0.85, 11.5, 0.8, [{"text": data.get("title", ""), "pt": 26, "color": "accent", "bold": True}], space_after=0)
    n = len(data["rows"])
    y = 1.85
    pitch = 4.6 / max(n, 1)
    if n <= 4:
        pitch = min(1.25, pitch)
    for i, r in enumerate(data["rows"], 1):
        _text(s, 0.9, y, 1.0, 0.6, [{"text": str(i), "pt": 22, "color": "accent", "bold": True}], space_after=0)
        _text(s, 2.1, y, 10.2, 0.6, [{"text": r["title"], "pt": 16, "color": "ink", "bold": True}], space_after=0)
        _text(s, 2.1, y + 0.6, 10.2, 0.6, [{"text": r.get("desc", ""), "pt": 14, "color": "muted"}], space_after=0)
        y += pitch

def L_compare(s, data, pal):
    _text(s, 0.5, 0.35, 12.3, 0.8, [{"text": data.get("title", ""), "pt": 26, "color": "accent", "bold": True}], space_after=0)
    for side, x0 in (("a", 0.5), ("b", 6.8)):
        d = data[side]
        _card(s, x0, 1.2, 6.0, 5.6, pal, rounded=True, border="muted")
        _text(s, x0 + 0.3, 1.45, 5.4, 0.7, [{"text": d.get("title", ""), "pt": 20, "color": "accent2" if side == "a" else "accent", "bold": True}], space_after=0)
        lines = []
        for r in d.get("rows", []):
            lines.append({"text": r["k"], "pt": 14, "color": "accent2", "bold": True})
            lines.append({"text": r["v"], "pt": 14, "color": "ink"})
        _text(s, x0 + 0.3, 2.25, 5.4, 4.2, lines, space_after=3)
        tk = data.get("takeaway_" + side) or data.get("takeaway")
        if tk:
            _text(s, x0 + 0.3, 5.9, 5.4, 0.8, [{"text": "★ " + tk, "pt": 14, "color": "accent", "bold": True}], space_after=0)

def L_process(s, data, pal):
    """Step-by-step: horizontal chevron-like cards with arrows."""
    n = len(data["steps"])
    w = (12.3 - 0.3 * (n - 1)) / n
    _text(s, 0.5, 0.35, 12.3, 0.8, [{"text": data.get("title", ""), "pt": 26, "color": "accent", "bold": True}], space_after=0)
    x = 0.5
    for i, st in enumerate(data["steps"]):
        _card(s, x, 1.4, w, 4.6, pal, rounded=True)
        lines = [{"text": f"STEP {i+1}", "pt": 13, "color": "accent2", "bold": True}]
        lines.append({"text": st.get("title", ""), "pt": 15, "color": "ink", "bold": True})
        lines.append({"text": st.get("desc", ""), "pt": 13, "color": "muted"})
        _text(s, x + 0.15, 1.7, w - 0.3, 4.0, lines, space_after=4)
        if i < n - 1:
            _text(s, x + w + 0.02, 3.5, 0.26, 0.5, [{"text": ">", "pt": 18, "color": "accent", "bold": True}], space_after=0)
        x += w + 0.3
    if data.get("footnote"):
        _text(s, 0.5, 6.5, 12.3, 0.5, [{"text": data["footnote"], "pt": 13, "color": "muted"}], space_after=0)

LAYOUTS = {"title": L_title, "bullets": L_bullets, "cards": L_cards, "two_panel": L_two_panel,
           "chart": L_chart, "numbered_rows": L_numbered_rows, "compare": L_compare, "process": L_process}

def make_deck(session, args):
    from pptx import Presentation
    spec = json.loads(args["spec"]) if isinstance(args.get("spec"), str) else args["spec"]
    path = args["path"]
    global _PAL
    theme = spec.get("theme", "dark")
    pal = dict(LIGHT if theme == "light" else DARK)
    _PAL = pal
    prs = Presentation()
    prs.slide_width = int(SW * EMU)
    prs.slide_height = int(SH * EMU)
    trans = spec.get("transition")
    for sd in spec["slides"]:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _bg(s, pal)
        if trans:
            _slide_in(s, trans)
        lay = LAYOUTS.get(sd.get("layout", "title"))
        lay(s, sd, pal)
    n_content = len(spec["slides"])
    prs.save(path)
    if n_content < 3:
        return f"WARN: saved {path} with only {n_content} slides; decks should have 3-8"
    return f"saved {path} ({n_content} slides)"


def check_deck(session, args):
    import importlib.util
    spec_path = os.path.join(os.path.dirname(__file__) or ".", "..", "..",
                            "deck-lab", "check_deck.py")
    spec_path = os.path.abspath(spec_path)
    if not os.path.exists(spec_path):
        spec_path = "/home/runner/work/manga-translator/manga-translator/deck-lab/check_deck.py"
    spec_ = importlib.util.spec_from_file_location("check_deck", spec_path)
    mod = importlib.util.module_from_spec(spec_)
    spec_.loader.exec_module(mod)
    png = args.get("png")
    import io, contextlib
    # Redirect stdout: check_deck.py's QC runner prints to stdout, and any
    # stray text on the MCP stdio transport breaks JSON-RPC framing.
    with contextlib.redirect_stdout(io.StringIO()):
        issues = mod.check(args["path"], png)
    return "=== CLEAN ===" if not issues else "ISSUES:\n" + "\n".join(issues)

INJECT = ["tools"]

def apply(ctx, config):
    t = ctx.get("tools")
    t.register(
        {"name": "make_deck",
         "description": "Build a .pptx from a JSON spec. spec: {theme: 'dark'|'light', slides: [{layout: title|bullets|cards|two_panel|chart|numbered_rows|compare|process, ...layout fields}]}. See .agents/skills/pptx/SKILL.md for layout field shapes. Writes to path.",
         "parameters": {"type": "object", "properties": {
             "spec": {"type": "string", "description": "JSON string of the deck spec"},
             "path": {"type": "string", "description": "output .pptx path"}},
             "required": ["spec", "path"]}},
        lambda s, a: make_deck(s, a), kind="edit")
    t.register(
        {"name": "check_deck_tool",
         "description": "QC a .pptx: text overflow, overlap, font floors, contrast, off-slide shapes. Optionally pass png prefix (pdftoppm output) for render checks. Returns CLEAN or issue list.",
         "parameters": {"type": "object", "properties": {
             "path": {"type": "string", "description": ".pptx file"},
             "png": {"type": "string", "description": "optional pdftoppm prefix, e.g. out/agents"}},
             "required": ["path"]}},
        lambda s, a: check_deck(s, a), kind="read")

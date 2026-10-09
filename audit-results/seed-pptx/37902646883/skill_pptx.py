"""A small, editorially opinionated PowerPoint toolkit.

Uses python-pptx. Coordinates are inches on a 13.333 x 7.5 16:9 canvas.
Prefer custom layouts and editable vector shapes over decorative templates.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

W, H = 13.333, 7.5
C = {"paper":"F6F5F1", "ink":"17211F", "muted":"59645F", "line":"D8D9D2", "accent":"146B5B", "pale":"E5EEE8", "white":"FFFFFF", "warm":"B96843", "dark":"132421"}
FONT = "Lato"

def color(hexstr):
    return RGBColor.from_string(hexstr.lstrip('#').upper())

def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.author = "Human-centered slide toolkit"
    prs.core_properties.subject = "Evidence-led, restrained presentation design"
    return prs

def rect(slide, x,y,w,h, fill, line=None, radius=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=color(fill)
    if line:
        sh.line.color.rgb=color(line); sh.line.width=Pt(0.8)
    else: sh.line.fill.background()
    if radius:
        try: sh.adjustments[0]=0.06
        except Exception: pass
    return sh

def text(slide, value, x,y,w,h, size=18, fill=None, bold=False, font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0, italic=False):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(margin)
    tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align; p.space_before=Pt(0); p.space_after=Pt(0)
    r=p.add_run(); r.text=str(value); r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic
    r.font.color.rgb=color(fill or C['ink'])
    return box

def label(slide, value, x,y,w=4, fill=None):
    return text(slide, value.upper(), x,y,w,.25, size=9, fill=fill or C['accent'], bold=True)

def base_slide(prs, title, section="FIELD NOTES", page=None, source=None, subtitle=None):
    slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=color(C['paper'])
    label(slide,section,.72,.42,5)
    text(slide,title,.72,.86,11.9,.94,size=28,bold=True)
    if subtitle: text(slide,subtitle,.74,1.91,11.6,.5,size=13,fill=C['muted'])
    rect(slide,.72,6.99,11.9,.012,C['line'])
    text(slide,source or "",.72,7.08,10.7,.2,size=8,fill=C['muted'])
    if page is not None: text(slide,f"{page:02d}",12.05,7.06,.55,.22,size=9,fill=C['muted'],align=PP_ALIGN.RIGHT)
    return slide

def title_slide(prs, title, subtitle, kicker="RESEARCH BRIEF", detail=""):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=color(C['dark'])
    rect(s,.78,.82,.11,.36,C['warm'])
    label(s,kicker,.78,1.42,5,"A9C9BA")
    text(s,title,.78,2.02,10.9,1.82,size=39,fill=C['white'],bold=True)
    text(s,subtitle,.8,4.2,9.8,.86,size=19,fill="CFDAD4")
    # restrained, aligned motif of gate layers; entirely within canvas
    for i,(yy,ww) in enumerate([(5.72,2.65),(6.04,3.1),(6.36,2.15)]):
        rect(s,9.55,yy,ww,.065, C['accent'] if i!=1 else C['warm'])
    text(s,detail,.8,6.85,9,.26,size=10,fill="A9B7B0")
    return s

def content_slide(prs,title,items,section="EXPLANATION",page=None,source=None,subtitle=None):
    """Single-column sequence: items are (short lead, explanatory sentence). Keep to 2-4."""
    s=base_slide(prs,title,section,page,source,subtitle)
    n=len(items); top=2.42 if subtitle else 2.18; gap=.17; row=(4.3-(n-1)*gap)/max(n,1)
    for i,(lead,body) in enumerate(items):
        y=top+i*(row+gap); rect(s,.74,y,.055,row-.04,C['accent'] if i==0 else C['line'])
        text(s,lead,.98,y+.04,3.0,.38,size=17,bold=True)
        text(s,body,4.12,y+.04,8.25,row-.08,size=15,fill=C['muted'])
    return s

def comparison_slide(prs,title,left_title,left_items,right_title,right_items,section="DECISION",page=None,source=None,subtitle=None):
    """Two-column comparison with matched criteria. Items may be (label, description)."""
    s=base_slide(prs,title,section,page,source,subtitle)
    top=2.25 if subtitle else 2.05
    for x,head,items,accent in [(.74,left_title,left_items,C['muted']),(6.82,right_title,right_items,C['accent'])]:
        rect(s,x,top,5.76,.62,C['dark'] if accent==C['accent'] else "E6E7E0")
        text(s,head,x+.2,top+.15,5.34,.3,size=16,fill=C['white'] if accent==C['accent'] else C['ink'],bold=True)
        for j,(lead,body) in enumerate(items):
            y=top+.88+j*1.02
            text(s,lead,x+.05,y,5.4,.28,size=13,bold=True,fill=C['accent'] if accent==C['accent'] else C['ink'])
            text(s,body,x+.05,y+.34,5.45,.54,size=11.5,fill=C['muted'])
    return s

def metric(slide,x,y,value,caption,accent=None):
    text(slide,value,x,y,3.45,.78,size=35,bold=True,fill=accent or C['accent'])
    text(slide,caption,x,y+.88,3.5,.75,size=13,fill=C['muted'])

def data_slide(prs,title,metrics,section="EVIDENCE",page=None,source=None,subtitle=None, takeaway=None):
    """Metric-led data slide. Metrics are (value, caption); avoid false precision."""
    s=base_slide(prs,title,section,page,source,subtitle)
    top=2.48 if subtitle else 2.27
    cols=len(metrics); total=11.85; gap=.28; cw=(total-gap*(cols-1))/cols
    for i,(value,caption) in enumerate(metrics):
        x=.74+i*(cw+gap)
        rect(s,x,top,cw,2.36,C['white'],C['line'])
        metric(s,x+.24,top+.36,value,caption, C['accent'] if i%2==0 else C['warm'])
    if takeaway:
        rect(s,.74,5.35,11.85,.92,C['pale'])
        text(s,takeaway,.98,5.59,11.3,.44,size=15,bold=True)
    return s

def save(prs,path):
    prs.save(path)
    return path

# Research behind defaults: python-pptx text and chart docs; WCAG 2.2 contrast;
# IBM Design color guidance; Reynolds/Presentation Zen narrative restraint;
# Tufte on quantitative integrity and avoiding chartjunk. See SKILL.md.
